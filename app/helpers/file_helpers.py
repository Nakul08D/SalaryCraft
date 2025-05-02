import io
import os
import smtplib
from dotenv import load_dotenv
from email.message import EmailMessage
import pandas as pd
from fastapi import FastAPI, UploadFile, File, HTTPException
from docx import Document
from typing import List
from .document_helpers import replace_placeholders, convert_docx_to_pdf
from .calculation_helpers import calculate_derived_fields


load_dotenv()

OUTPUT_DIR = "generated_payslips"
os.makedirs(OUTPUT_DIR, exist_ok=True)


async def save_and_generate_pdf(csv_file: UploadFile) -> List[str]:
    status_messages = []

    try:
        contents = await csv_file.read()
        df = pd.read_csv(io.StringIO(contents.decode('utf-8')))
        df = df.replace({pd.NA: None, '': None})
        records = df.to_dict(orient='records')
    except Exception as e:
        return [f"Failed to process CSV file: {str(e)}"]

    for i, record in enumerate(records):
        try:
            try:
                record = calculate_derived_fields(record)
            except Exception as e:
                status_messages.append(f"Error calculating derived fields for record {i+1}: {str(e)}")
                continue

            try:
                doc = Document("app/static/Dummy.docx")
                doc = replace_placeholders(doc, record)
            except Exception as e:
                status_messages.append(f"{record.get('Name', f'Employee_{i+1}')}: Error processing DOCX template: {str(e)}")
                continue

            employee_name = record.get('Name', f"Employee_{i+1}").replace(" ", "_")
            month = record.get('Month', 'Payslip').replace(" ", "_")
            docx_filename = f"{employee_name}_Salary_Slip_{month}.docx"

            send_mail_flag = record.get('Send_Mail', '').strip().lower()
            email = record.get("Email")

            target_dir = "not_sent_payslips" if send_mail_flag == 'no' else "generated_payslips"

            try:
                os.makedirs(target_dir, exist_ok=True)
                docx_path = os.path.join(target_dir, docx_filename)
                doc.save(docx_path)
            except Exception as e:
                status_messages.append(f"{employee_name}: Error saving DOCX file: {str(e)}")
                continue

            try:
                pdf_path = docx_path.replace(".docx", ".pdf")
                convert_docx_to_pdf(docx_path, pdf_path)

                if not os.path.exists(pdf_path):
                    status_messages.append(f"{employee_name}: PDF conversion failed.")
                    continue

                os.remove(docx_path)
            except Exception as e:
                status_messages.append(f"{employee_name}: Error during PDF conversion or DOCX cleanup: {str(e)}")
                continue

            if send_mail_flag != 'no':
                if email:
                    try:
                        mail_body = (
                            f"Dear {record.get('Name', '')},\n\n"
                            f"We hope this email finds you well.\n\n"
                            f"Please find your salary slip for the month of {record.get('Month', '')}.\n"
                            f"Kindly review it at your earliest convenience and let us know if you have any questions.\n\n"
                            f"Best regards,\n"
                            f"Cubexo HR Team"
                        )
                        subject = f"Your Payslip - {record.get('Month', '')}"
                        send_email_with_pdf(email, subject, mail_body, pdf_path)
                        status_messages.append(f"{employee_name}: Payslip emailed to the employee.")
                    except Exception as e:
                        status_messages.append(f"{employee_name}: Failed to send email: {str(e)}")
                else:
                    status_messages.append(f"{employee_name}: Email not provided. Payslip not sent.")
            else:
                status_messages.append(f"{employee_name}: Payslip generated, but email not sent as requested.")

        except Exception as e:
            status_messages.append(f"{employee_name}: Unexpected error: {str(e)}")

    return status_messages


def send_email_with_pdf(receiver_email, subject, body, pdf_path):
    try:
        sender_email = os.getenv("SENDER_EMAIL")
        sender_password = os.getenv("SENDER_PASSWORD")
        
        msg = EmailMessage()
        msg['Subject'] = subject
        msg['From'] = sender_email
        msg['To'] = receiver_email
        msg.set_content(body)

        with open(pdf_path, 'rb') as f:
            file_data = f.read()
            file_name = os.path.basename(pdf_path)

        msg.add_attachment(file_data, maintype='application', subtype='pdf', filename=file_name)

        with smtplib.SMTP('smtp.gmail.com', 587) as smtp:
            smtp.starttls()
            smtp.login(sender_email, sender_password)
            smtp.send_message(msg)

        print(f"Payslip sent to {receiver_email}")
    except Exception as e:
        print(f"Failed to send to {receiver_email}: {e}")
